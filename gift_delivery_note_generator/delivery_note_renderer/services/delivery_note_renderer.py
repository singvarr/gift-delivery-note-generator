from __future__ import annotations

import copy
import shutil
from typing import TYPE_CHECKING
from logging import getLogger
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.table import _Cell
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from gift_delivery_note_generator.delivery_note_renderer.constants.document_settings import (
    FONT_NAME,
    FONT_SIZE_PT,
    ITALIC_TAG,
)

from ..models.delivery_note import DeliveryNote
from ..settings.utils.build_delivery_note_name import build_delivery_note_name

if TYPE_CHECKING:
    from gift_delivery_note_generator.app import Config


class DeliveryNoteRenderer:
    def __init__(self, delivery_note: DeliveryNote, config: Config) -> None:
        self._delivery_note = delivery_note
        self._config = config

        self._logger = getLogger(__name__)

    @property
    def _destination_path(self) -> Path:
        document_name = build_delivery_note_name(self._delivery_note, self._config)

        return self._config.output_path / f"{document_name}.docx"

    @property
    def _context(self) -> dict[str, str | int]:
        return {
            "{{TOTAL_GIFTS}}": self._delivery_note.total_gifts,
            "{{TOTAL_HUMANIZED}}": self._delivery_note.humanized_total_gifts,
            "{{ISSUE_DATE}}": self._delivery_note.issue_date,
            "{{RECIPIENT}}": self._delivery_note.store_id,
        }

    def _create_doc_from_template(self) -> DocumentType:
        shutil.copy(self._config.template_path, self._destination_path)

        return Document(str(self._destination_path))

    def _replace_tags_by_values(self, document: DocumentType) -> None:
        paragraphs = list(document.paragraphs)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)

        for paragraph in paragraphs:
            self._replace_tags_in_paragraph(paragraph, self._context)

    def _replace_tags_in_paragraph(
        self,
        paragraph: Paragraph,
        context: dict[str, str | int],
    ) -> None:
        full_text = "".join(run.text for run in paragraph.runs)

        if not full_text or not paragraph.runs:
            return

        replaced_text = full_text
        for tag, value in context.items():
            if tag != ITALIC_TAG:
                replaced_text = replaced_text.replace(tag, str(value))

        if replaced_text == full_text and ITALIC_TAG not in full_text:
            return

        for run in paragraph.runs[1:]:
            run.text = ""

        before, tag, after = replaced_text.partition(ITALIC_TAG)
        first_run = paragraph.runs[0]
        first_run.text = before

        if tag:
            self._add_run_after(first_run, after)
            italic_run = self._add_run_after(first_run, str(context[ITALIC_TAG]))
            italic_run.font.italic = True

    def _add_run_after(self, run: Run, text: str) -> Run:
        new_element = copy.deepcopy(run._r)
        run._r.addnext(new_element)

        new_run = Run(new_element, run._parent)
        new_run.text = text

        return new_run

    def _print_gift_table(self, document: DocumentType) -> None:
        table = document.tables[0]
        footer_row = table.rows[-1]

        for gift_entry in self._delivery_note.gifts:
            row = table.add_row()

            row.cells[0].text = gift_entry.order_details
            row.cells[1].text = gift_entry.recipient
            row.cells[2].text = gift_entry.store_id

            for index, cell in enumerate(row.cells):
                self._format_cell(cell, bold=index == 0)

            footer_row._tr.addprevious(row._tr)

    def _format_cell(self, cell: _Cell, bold: bool = False) -> None:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in paragraph.runs:
                run.font.name = FONT_NAME
                run.font.size = Pt(FONT_SIZE_PT)
                run.font.bold = bold

    def run(self) -> None:
        document = self._create_doc_from_template()

        self._print_gift_table(document=document)
        self._replace_tags_by_values(document=document)

        document.save(str(self._destination_path))
        self._logger.info(f"Created delivery note {self._destination_path.name}")

from logging import getLogger
import shutil

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from gift_delivery_note_generator.models.delivery_note import DeliveryNote
from gift_delivery_note_generator.constants.paths import BASE_TEMPLATE_PATH, OUTPUT_PATH
from gift_delivery_note_generator.store_config.utils.build_delivery_note_name import (
    build_delivery_note_name,
)


class DeliveryNoteRenderer:
    def __init__(self, delivery_note: DeliveryNote):
        self._delivery_note = delivery_note

        self._logger = getLogger(__name__)

    @property
    def _destination_path(self):
        document_name = build_delivery_note_name(self._delivery_note)

        return OUTPUT_PATH / f"{document_name}.docx"

    @property
    def _context(self):
        return {
            "{{TOTAL_GIFTS}}": self._delivery_note.total_gifts,
            "{{TOTAL_HUMANIZED}}": self._delivery_note.humanized_total_gifts,
            "{{ISSUE_DATE}}": self._delivery_note.issue_date,
            "{{RECIPIENT}}": self._delivery_note.store_id,
        }

    def _create_doc_from_template(self) -> Document:
        shutil.copy(BASE_TEMPLATE_PATH, self._destination_path)

        return Document(self._destination_path)

    def _replace_tags_by_values(self, document: Document):
        paragraphs = list(document.paragraphs)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)

        for paragraph in paragraphs:
            self._replace_tags_in_paragraph(paragraph, self._context)

    def _replace_tags_in_paragraph(self, paragraph, context: dict[str, str]):
        full_text = "".join(run.text for run in paragraph.runs)

        if not full_text or not paragraph.runs:
            return

        replaced_text = full_text
        for tag, value in context.items():
            replaced_text = replaced_text.replace(tag, str(value))

        if replaced_text == full_text:
            return

        paragraph.runs[0].text = replaced_text
        for run in paragraph.runs[1:]:
            run.text = ""

    def _print_gift_table(self, document: Document):
        table = document.tables[0]
        footer_row = table.rows[-1]

        for gift_entry in self._delivery_note.gifts:
            row = table.add_row()

            row.cells[0].text = gift_entry.order_details
            row.cells[1].text = gift_entry.recipient
            row.cells[2].text = gift_entry.store_id

            for index, cell in enumerate(row.cells):
                self._format_cell(cell, bold=index == 0)

            footer_row._tr.addprevious(row._tr)

    def _format_cell(self, cell, bold: bool = False):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.font.bold = bold

    def run(self):
        document = self._create_doc_from_template()

        self._print_gift_table(document=document)
        self._replace_tags_by_values(document=document)

        document.save(self._destination_path)

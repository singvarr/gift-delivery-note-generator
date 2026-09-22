
from typing import Type, TypeVar

import openpyxl

from .constants.keyword_separator import KEYWORD_SEPARATOR
from .models.column_mapping import ColumnMapping
from .models.table_settings import TableSettings

T = TypeVar("T")


class ExcelTableParser:
    def __init__(
        self,
        table_settings: TableSettings,
        dataclass_type: Type[T],
        mappings: tuple[ColumnMapping, ...]
    ):
        self._table_settings = table_settings
        self._dataclass_type = dataclass_type
        self._mappings = mappings

    def run(self) -> list[T]:
        workbook = openpyxl.load_workbook(self._table_settings.path, data_only=True)

        table = None
        sheet = None

        for worksheet in workbook.worksheets:
            if self._table_settings.table_name in worksheet.tables:
                table = worksheet.tables[self._table_settings.table_name]
                sheet = worksheet
                break

        if table is None:
            raise Exception(
                f"Table '{self._table_settings.table_name}' not found in "
                f'{self._table_settings.path}'
            )

        data_range = list(sheet[table.ref])
        header_row, *data_rows = data_range

        headers = [cell.value for cell in header_row]

        result: list[T] = []

        for row in data_rows:
            data = {}

            for mapping in self._mappings:
                index = headers.index(mapping.column)

                if mapping.is_keyword_field:
                    value = row[index].value.split(KEYWORD_SEPARATOR)
                else:
                    value = row[index].value

                data[mapping.field] = value

            result.append(self._dataclass_type(**data))

        return result
